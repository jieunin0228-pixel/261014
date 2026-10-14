"""3·4·5단계 시험. Claude 호출(ai.call_json)은 가짜 응답으로 바꿉니다(실제 응답 품질은 시험하지 않음)."""
import io
import os
import re
import shutil
import sys
import tempfile
import time
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
TMP = tempfile.mkdtemp(prefix="terms_test3_")
os.environ["TERMS_DATA_DIR"] = os.path.join(TMP, "data")

from openpyxl import Workbook, load_workbook  # noqa: E402

from converter import ai, cards, compare, ecs, export, pipeline  # noqa: E402
from tests import make_fixture  # noqa: E402

pipeline.ROOT = os.environ["TERMS_DATA_DIR"]
RIDER = "R(무)테스트암진단특약(갱신형)"


def convert(path, name):
    up = os.path.join(TMP, name + ".up")
    shutil.copy(path, up)
    job = pipeline.start_conversion(up, name)
    while not job["done"]:
        time.sleep(0.2)
    assert not job["error"], job["error"]
    return job["doc_id"]


def prep(doc):
    rider = next(c for c in pipeline.read_json(doc, "coverage_map.json")["coverages"] if c["name"] == RIDER)
    pipeline.add_edit(doc, {"action": "coverage_range", "cid": rider["cid"], "start": 6, "end": 8})
    pipeline.set_approval(doc, True, "시험")
    return rider["cid"]


def handle(text, label):
    return re.search(r"\[([A-Z]\d+)\] \(\d+쪽\) " + re.escape(label), text).group(1)


def fake_card(system, user, schema, effort="high", max_tokens=32000):
    pay = handle(user, "제2-2조")
    tbl = re.search(r"\[(T\d+)\] 표", user).group(1)
    return {
        "items": [
            {"field": "지급사유", "content": "암보장개시일 이후 암으로 진단확정되었을 때 지급합니다.", "status": "명시",
             "evidence": [{"ref": pay, "quote": "암보장개시일 이후에 암으로 진단확정되었을 때"}]},
            {"field": "감액기간과 감액비율", "content": "1년 미만 50% 지급", "status": "명시",
             "evidence": [{"ref": pay, "quote": "최초계약의 계약일부터 1년 미만 진단확정시 50%를 지급합니다"}]},
            {"field": "면책기간", "content": "90일 면책", "status": "명시",
             "evidence": [{"ref": pay, "quote": "계약일부터 90일 이내에는 지급하지 않습니다"}]},   # 원문에 없음 → 격리
            {"field": "지급횟수와 한도", "content": "약관에 지급횟수 규정이 없습니다.", "status": "약관에 명시 없음", "evidence": []},
        ],
        "cautions": [
            {"topic": "경과기간·연령·성별별 지급률 차이", "content": "1년 미만 진단 시 50% 감액",
             "needs_interpretation": True, "interpretations": ["진단확정일 기준", "청구일 기준"],
             "evidence": [{"ref": pay, "quote": "1년 미만 진단확정시 50%를 지급합니다"}]},
        ],
        "tables": [{"ref": tbl, "use": "지급금액"}],
    }, {"model": "fake"}


def fake_ecs(system, user, schema, effort="high", max_tokens=32000):
    pay = handle(user, "제2-2조")
    tbl = re.search(r"\[(T\d+)\] 표", user).group(1)
    return {"rows": [{
        "label": "암진단자금 · 진단 · 질병", "split_reason": "단일 급부", "split_evidence": [],
        "fields": [
            {"field": "청구사유", "code": "S15", "name": "진단", "status": "값", "note": "",
             "evidence": [{"ref": pay, "quote": "암으로 진단확정되었을 때"}]},
            {"field": "분자", "code": "", "name": "100", "status": "값", "note": "",
             "evidence": [{"ref": tbl, "quote": "특약보험가입금액의"}]},
            {"field": "사인그룹", "code": "", "name": "암", "status": "값", "note": "",
             "evidence": [{"ref": pay, "quote": "없는 문장입니다"}]},
            {"field": "갱신가능불가", "code": "", "name": "", "status": "사용자 입력 필요", "note": "", "evidence": []},
        ]}]}, {"model": "fake"}


class Stage345(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.new = convert(make_fixture.build(os.path.join(TMP, "new.pdf")), "new.pdf")
        cls.old = convert(make_fixture.build(os.path.join(TMP, "old.pdf"), old=True), "old.pdf")
        cls.cid = prep(cls.new)
        cls.old_cid = prep(cls.old)
        os.environ["ANTHROPIC_API_KEY"] = "test"
        cls.orig = ai.call_json

    @classmethod
    def tearDownClass(cls):
        ai.call_json = cls.orig
        os.environ.pop("ANTHROPIC_API_KEY", None)

    def test_1_bundle_rule6(self):
        b = ai.Bundle(self.new, self.cid)
        common, user = b.build()
        self.assertIn("제1-2조【용어의 정의】", common)               # (b) 특약 공통 조항
        self.assertIn("제2-2조【보험금의 지급사유】", user)            # (a) 담보 구간
        self.assertIn("「의료법」", user)                              # (d) 인용조문
        self.assertIn("주계약 약관 준용 표시", user)                    # (c) 준용 표시
        self.assertNotIn("제3조【보험금을 지급하지 않는 사유】", user)  # 참조하지 않은 보통약관 조문은 보내지 않음

    def test_2_cards(self):
        ai.call_json = fake_card
        cards.run(self.new, [self.cid])
        card = cards.load(self.new)[self.cid]
        fields = {i["field"]: i for i in card["items"]}
        self.assertEqual(fields["지급사유"]["check"], "근거 확인")
        self.assertEqual(fields["지급사유"]["evidence"][0]["id"], RIDER + "/제2-2조")
        self.assertEqual(fields["지급횟수와 한도"]["check"], "약관에 명시 없음")
        self.assertEqual([i["field"] for i in card["isolated"]], ["면책기간"])   # 근거 불일치 격리
        self.assertEqual(fields["면책기간"]["check"], "근거 불일치로 격리")        # 격리 후 빈 항목 표시
        self.assertEqual(fields["보장개시일"]["check"], "AI 미응답")
        self.assertTrue(card["cautions"][0]["needs_interpretation"])
        self.assertTrue(card["tables"])
        cards.set_review(self.new, self.cid, fields["지급사유"]["id"], "검토완료")

    def test_3_ecs_without_and_with_master(self):
        ai.call_json = fake_card
        cards.run(self.new, [self.cid])
        ai.call_json = fake_ecs
        if os.path.exists(ecs.master_path()):
            os.remove(ecs.master_path())
        ecs.run(self.new, [self.cid])
        row = ecs.load(self.new)[self.cid]["rows"][0]
        f = row["fields"]
        self.assertEqual((f["청구사유"]["code"], f["청구사유"]["name"]), ("", "진단"))
        self.assertIn("코드 확인 필요", f["청구사유"]["flags"])
        self.assertEqual(f["분자"]["name"], "100")
        self.assertEqual(f["사인그룹"]["status"], "약관에 명시 없음")   # 근거 불일치 → 비움
        self.assertEqual(row["isolated"][0]["field"], "사인그룹")
        self.assertEqual(f["상품코드"]["status"], "사용자 입력 필요")
        self.assertEqual(f["계약상태처리"]["status"], "약관에 명시 없음")
        self.assertEqual(list(f), ecs.FIELD_KEYS)
        # 코드마스터 적용
        wb = Workbook()
        ws = wb.active
        ws.append(["필드명", "코드", "명칭"])
        ws.append(["청구사유", "S15", "진단"])
        bio = io.BytesIO()
        wb.save(bio)
        ecs.import_master(bio.getvalue())
        ecs.run(self.new, [self.cid])
        f = ecs.load(self.new)[self.cid]["rows"][0]["fields"]
        self.assertEqual(f["청구사유"]["code"], "S15")
        ecs.set_product(self.new, {"상품코드": "439021", "상품명": "시험상품", "판매개시일자": "20260501"})
        v = ecs.view_rows(self.new, self.cid)
        self.assertEqual(v["rows"][0]["fields"]["상품코드"]["name"], "439021")
        ecs.edit_cell(self.new, self.cid, "R1", "계약상태처리", "", "소멸")
        self.assertEqual(ecs.load(self.new)[self.cid]["rows"][0]["fields"]["계약상태처리"]["name"], "소멸")
        x = export.excel(self.new, [self.cid])
        names = load_workbook(io.BytesIO(x)).sheetnames
        self.assertIn("결과카드", names)
        self.assertIn("ECS초안", names)
        os.remove(ecs.master_path())

    def test_4_requires_approval(self):
        pipeline.set_approval(self.new, False)
        try:
            with self.assertRaises(PermissionError):
                cards.run(self.new, [self.cid])
            with self.assertRaises(PermissionError):
                compare.run(self.new, self.old)
        finally:
            pipeline.set_approval(self.new, True)

    def test_5_compare(self):
        data = compare.run(self.new, self.old)
        p = next(p for p in data["pairs"] if p["new"] == self.cid)
        self.assertEqual(p["old"], self.old_cid)
        self.assertEqual(p["status"], "변경")
        self.assertEqual(p["opinion"], "의미 변경 가능성 높음")
        art = next(a for a in p["articles"] if a["key"] == "제2-2조")
        self.assertIn(["-", next(t for op, t in art["ops"] if op == "-")], art["ops"])
        self.assertTrue(any("30%" in t for op, t in art["ops"] if op == "-"))
        self.assertTrue(any("50%" in t for op, t in art["ops"] if op == "+"))
        self.assertIn("청구사유", p["fields"])
        main = next(p for p in data["pairs"] if p["new_name"] and "건강보험" in p["new_name"])
        self.assertEqual(main["status"], "동일")
        # 짝 바꾸기: 짝 없음으로 두면 신규/삭제로 나뉩니다.
        data = compare.set_match(self.new, self.cid, "")
        self.assertEqual(next(p for p in data["pairs"] if p["new"] == self.cid)["status"], "신규")
        self.assertTrue(any(p["status"] == "삭제" and p["old"] == self.old_cid for p in data["pairs"]))
        data = compare.set_match(self.new, self.cid, self.old_cid)
        x = compare.excel(self.new)
        self.assertEqual(load_workbook(io.BytesIO(x)).sheetnames, ["비교요약", "조문변경", "ECS 수정전후"])

    def test_6_locate(self):
        import app as appmod
        loc = appmod.locate(self.new, node=RIDER + "/제2-2조", quote="1년 미만 진단확정시 50%를 지급합니다")
        self.assertTrue(loc["found"])
        self.assertEqual(loc["page"], 6)
        raw = pipeline.read_json(self.new, "raw_pages.json")["pages"][5]["lines"]
        self.assertIn("50%", raw[loc["lines"][-1]]["t"])


if __name__ == "__main__":
    unittest.main()
