"""5단계: 구약관 대비 비교.

- 비교 기준은 구약관 PDF입니다. 구약관도 1단계 변환과 승인을 거친 문서만 고를 수 있습니다.
- 담보 매칭: 담보명 유사도(상품 꼬리표를 뗀 이름)와 목차상 위치로 자동 제안하고, 사용자가 바꿀 수 있습니다.
  짝이 없으면 신규/삭제 후보로 표시합니다. 변경구분은 신규·변경·삭제·동일입니다.
- 조문·별표 단위로 구 원문과 신 원문을 대조하고(공백 차이는 무시), 바뀐 부분을 표시합니다.
- 의미 변경 여부는 확정하지 않고 "의미 변경 가능성 높음 / 표현 변경으로 보임" 의견만 붙입니다(규칙 기반, AI 미사용).
- 변경 건마다 "ECS 룰 수정 필요 필드"를 붙입니다. 양쪽에 4단계 ECS 초안이 있으면 필드별 차이로, 없으면 조문 종류로 추정합니다.
"""
from __future__ import annotations

import datetime as dt
import difflib
import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from . import ai, ecs, pipeline
from .classify import short_name
from .extract import norm_text

TYPES = {"보통약관", "특약공통", "개별특약", "납입면제류", "서비스·인수·전환류", "미분류"}
MEANING = re.compile(r"\d|%|일|년|개월|회|세|원|배|다만|제외|않|한하여|한함|이내|이상|이하|미만|초과|포함|최초|1회|감액|면책|지급하지")

FIELD_BY_TOPIC = [
    (r"지급사유|납입면제사유", ["청구사유", "사고유형", "사인그룹"] + [f"급부결정조건{i}" for i in range(1, 7)]),
    (r"세부규정", [f"급부결정조건{i}" for i in range(1, 7)] + [f"계산식결정조건{i}" for i in range(1, 7)] + ["계산식"]),
    (r"지급하지|면제하지|면책", ["사인그룹"] + [f"급부결정조건{i}" for i in range(1, 7)]),
    (r"보장개시", [f"급부결정조건{i}" for i in range(1, 7)]),
    (r"보험기간|갱신", ["갱신가능불가"]),
    (r"소멸", ["계약상태처리"]),
    (r"납입면제|보험료의납입", ["납면처리"]),
    (r"지급기준표", ["분자", "분모", "계산식", "상세계산식1", "금액유형1"]),
    (r"정의|분류표|진단확정|수술|입원", ["사인그룹"] + [f"급부결정조건{i}" for i in range(1, 7)]),
]


def now():
    return dt.datetime.now().isoformat(timespec="seconds")


def load(doc_id):
    return pipeline.read_json(doc_id, "compare.json")


def save(doc_id, data):
    pipeline.write_json(doc_id, "compare.json", data)


def _covs(doc_id):
    return [c for c in pipeline.read_json(doc_id, "coverage_map.json")["coverages"] if c["type"] in TYPES]


def _key(name):
    s = norm_text(short_name(name))
    return re.sub(r"\d+(\.\d+)+", "", s)   # 상품 버전 번호(3.10.5 등) 차이는 무시


def match(new_doc, old_doc, overrides=None):
    overrides = overrides or {}
    new, old = _covs(new_doc), _covs(old_doc)
    pairs, used_old = {}, set()
    for n in new:
        if n["cid"] in overrides:
            o = overrides[n["cid"]]
            pairs[n["cid"]] = (o or None, 1.0 if o else 0.0, "사용자")
            if o:
                used_old.add(o)
    cands = []
    for i, n in enumerate(new):
        if n["cid"] in pairs:
            continue
        for j, o in enumerate(old):
            if o["cid"] in used_old:
                continue
            if n["type"] != o["type"] and "미분류" not in (n["type"], o["type"]):
                continue
            nm = difflib.SequenceMatcher(None, _key(n["name"]), _key(o["name"])).ratio()
            pos = 1 - abs(i / max(1, len(new) - 1) - j / max(1, len(old) - 1))
            cands.append((0.85 * nm + 0.15 * pos, nm, n["cid"], o["cid"]))
    for score, nm, nc, oc in sorted(cands, reverse=True):
        if nc in pairs or oc in used_old or nm < 0.6:
            continue
        pairs[nc] = (oc, round(score, 3), "자동")
        used_old.add(oc)
    out = []
    for n in new:
        oc, score, how = pairs.get(n["cid"], (None, 0.0, "자동"))
        out.append({"new": n["cid"], "old": oc, "score": score, "how": how})
    for o in old:
        if o["cid"] not in used_old:
            out.append({"new": None, "old": o["cid"], "score": 0.0, "how": "자동"})
    return out


def _articles(doc_id, cid, nodes, orders):
    arts = []
    for nid in orders.get(cid, []):
        n = nodes[nid]
        if n["parent"] is None and n["kind"] in ("article", "annex"):
            arts.append(n)
    return arts


def _tokens(text):
    return re.findall(r"\S+", ai.join_lines(text))


def _diff_ops(a, b):
    """낱말 단위 차이: [["=", 글], ["-", 지운 글], ["+", 더한 글]]"""
    ta, tb = _tokens(a), _tokens(b)
    sm = difflib.SequenceMatcher(None, ta, tb, autojunk=False)
    ops = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            ops.append(["=", " ".join(ta[i1:i2])])
        else:
            if i2 > i1:
                ops.append(["-", " ".join(ta[i1:i2])])
            if j2 > j1:
                ops.append(["+", " ".join(tb[j1:j2])])
    return ops


def _opinion(ops):
    changed = " ".join(t for op, t in ops if op != "=")
    return "의미 변경 가능성 높음" if MEANING.search(changed) else "표현 변경으로 보임"


def _fields_for(title):
    t = norm_text(title)
    out = []
    for rx, fs in FIELD_BY_TOPIC:
        if re.search(rx, t):
            out += [f for f in fs if f not in out]
    return out


def _table_text(tables, owner):
    return "\n".join(ai.table_text(t) for t in tables if t["owner"] == owner)


def compare_pair(new_doc, old_doc, nc, oc, N, O):
    """N, O: 각 문서의 (nodes, orders, tables) 묶음."""
    na = _articles(new_doc, nc, N["nodes"], N["orders"]) if nc else []
    oa = _articles(old_doc, oc, O["nodes"], O["orders"]) if oc else []
    left = {a["key"]: a for a in oa}
    used = set()
    rows = []
    for a in na:
        b = left.get(a["key"])
        if b and difflib.SequenceMatcher(None, norm_text(a.get("title", "")), norm_text(b.get("title", ""))).ratio() < 0.5:
            b = None
        if b is None:  # 번호가 바뀐 경우 제목으로 찾습니다.
            b = next((x for x in oa if x["key"] not in used and norm_text(x.get("title", "")) == norm_text(a.get("title", ""))), None)
        if b is not None:
            used.add(b["key"])
        rows.append((b, a))
    rows += [(b, None) for b in oa if b["key"] not in used]
    arts = []
    for b, a in rows:
        nt = ai.full_text(N["nodes"], a["id"]) + ("\n[표]\n" + _table_text(N["tables"], a["id"]) if a and a["kind"] == "annex" else "") if a else ""
        ot = ai.full_text(O["nodes"], b["id"]) + ("\n[표]\n" + _table_text(O["tables"], b["id"]) if b and b["kind"] == "annex" else "") if b else ""
        if a and b:
            if norm_text(nt) == norm_text(ot):
                st, ops, op = "동일", [], ""
            else:
                ops = _diff_ops(ot, nt)
                st, op = "변경", _opinion(ops)
        elif a:
            st, ops, op = "신규", [["+", ai.join_lines(nt)]], "의미 변경 가능성 높음"
        else:
            st, ops, op = "삭제", [["-", ai.join_lines(ot)]], "의미 변경 가능성 높음"
        title = (a or b).get("title", "")
        arts.append({"key": (a or b)["key"], "title": title, "status": st, "opinion": op, "ops": ops,
                     "new_id": a["id"] if a else None, "old_id": b["id"] if b else None,
                     "new_page": a["page_start"] if a else None, "old_page": b["page_start"] if b else None,
                     "old_key": b["key"] if b else None,
                     "fields": _fields_for(title) if st != "동일" else []})
    return arts


def run(new_doc, old_doc, user="", overrides=None):
    for d in (new_doc, old_doc):
        if not pipeline.is_approved(d):
            raise PermissionError("비교하려면 신약관과 구약관 모두 1단계 변환이 승인되어 있어야 합니다.")
    if new_doc == old_doc:
        raise ValueError("같은 문서끼리는 비교할 수 없습니다.")
    prev = load(new_doc) or {}
    if overrides is None:
        overrides = prev.get("overrides", {}) if prev.get("old_doc") == old_doc else {}

    def pack(d):
        t = pipeline.read_json(d, "terms.json")
        return {"nodes": t["nodes"], "orders": t["orders"], "tables": pipeline.read_json(d, "tables.json")["tables"],
                "covs": {c["cid"]: c for c in pipeline.read_json(d, "coverage_map.json")["coverages"]}}
    N, O = pack(new_doc), pack(old_doc)
    new_ecs, old_ecs = ecs.load(new_doc), ecs.load(old_doc)
    pairs = []
    for m in match(new_doc, old_doc, overrides):
        nc, oc = m["new"], m["old"]
        arts = compare_pair(new_doc, old_doc, nc, oc, N, O)
        if nc and not oc:
            status = "신규"
        elif oc and not nc:
            status = "삭제"
        else:
            status = "동일" if all(a["status"] == "동일" for a in arts) else "변경"
        changed = [a for a in arts if a["status"] != "동일"]
        opinion = "" if status == "동일" else ("의미 변경 가능성 높음" if any(a["opinion"] == "의미 변경 가능성 높음" for a in changed) else "표현 변경으로 보임")
        fields, field_src = [], "조문 종류로 추정"
        if nc in new_ecs and oc in old_ecs:
            nv = ecs.view_rows(new_doc, nc, old_ecs[oc]["rows"])
            for r in nv["rows"]:
                for k, cell in r["fields"].items():
                    bc = cell.get("before") or {}
                    if (bc.get("code"), bc.get("name")) != (cell["code"], cell["name"]) and k not in fields:
                        fields.append(k)
            field_src = "양쪽 4단계 ECS 초안의 필드별 차이"
            if not fields and status == "변경":
                for a in changed:
                    fields += [f for f in a["fields"] if f not in fields]
                field_src = "ECS 초안 값은 같지만 조문이 바뀜 — 조문 종류로 추정(초안 재생성 권장)"
        else:
            for a in changed:
                fields += [f for f in a["fields"] if f not in fields]
        pairs.append({"new": nc, "old": oc, "new_name": N["covs"][nc]["name"] if nc else None,
                      "old_name": O["covs"][oc]["name"] if oc else None,
                      "new_range": [N["covs"][nc]["start"], N["covs"][nc]["end"]] if nc else None,
                      "old_range": [O["covs"][oc]["start"], O["covs"][oc]["end"]] if oc else None,
                      "score": m["score"], "how": m["how"], "status": status, "opinion": opinion,
                      "fields": fields, "field_source": field_src, "articles": arts})
    meta_old = pipeline.read_json(old_doc, "meta.json") or {}
    data = {"old_doc": old_doc, "old_filename": meta_old.get("filename"), "generated_at": now(), "overrides": overrides,
            "pairs": pairs,
            "counts": {s: sum(1 for p in pairs if p["status"] == s) for s in ("신규", "변경", "삭제", "동일")}}
    save(new_doc, data)
    pipeline._history(new_doc, "5단계 구약관 비교", {"old_doc": old_doc, "counts": data["counts"]}, user)
    return data


def set_match(new_doc, new_cid, old_cid, user=""):
    data = load(new_doc)
    if not data:
        raise KeyError("비교 결과가 없습니다.")
    ov = data.get("overrides", {})
    ov[new_cid] = old_cid or ""
    pipeline._history(new_doc, "5단계 담보 짝 변경", {"new": new_cid, "old": old_cid}, user)
    return run(new_doc, data["old_doc"], user, ov)


def excel(new_doc, only_changed=True):
    data = load(new_doc)
    if not data:
        raise KeyError("비교 결과가 없습니다.")
    wb = Workbook()
    ws = wb.active
    ws.title = "비교요약"
    head = ["변경구분", "의견", "신 담보ID", "신 담보명", "신 쪽", "구 담보ID", "구 담보명", "구 쪽", "매칭", "ECS 룰 수정 필요 필드", "필드 판단 근거"]
    ws.append(head)
    for p in data["pairs"]:
        if only_changed and p["status"] == "동일":
            continue
        ws.append([p["status"], p["opinion"], p["new"], p["new_name"], "-".join(map(str, p["new_range"] or [])),
                   p["old"], p["old_name"], "-".join(map(str, p["old_range"] or [])),
                   f"{p['how']} {p['score']}", ", ".join(p["fields"]), p["field_source"]])
    ws2 = wb.create_sheet("조문변경")
    ws2.append(["신 담보ID", "조문", "제목", "변경구분", "의견", "구 쪽", "신 쪽", "구 원문(바뀐 부분)", "신 원문(바뀐 부분)", "ECS 수정 필요 필드"])
    for p in data["pairs"]:
        for a in p["articles"]:
            if a["status"] == "동일":
                continue
            old = " ".join(t if op == "=" else f"[[{t}]]" for op, t in a["ops"] if op in ("=", "-"))
            new = " ".join(t if op == "=" else f"[[{t}]]" for op, t in a["ops"] if op in ("=", "+"))
            ws2.append([p["new"] or p["old"], a["key"], a["title"], a["status"], a["opinion"], a["old_page"], a["new_page"],
                        old, new, ", ".join(a["fields"])])
    for row in ws2.iter_rows(min_row=2, min_col=8, max_col=9):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws3 = wb.create_sheet("ECS 수정전후")
    ws3.append(["신 담보ID", "담보명", "룰 행", "필드", "수정 전 코드", "수정 전 명칭", "수정 후 코드", "수정 후 명칭", "변경", "상태", "근거 쪽(신)"])
    fill = PatternFill("solid", fgColor="FFF2CC")
    for p in data["pairs"]:
        if only_changed and p["status"] == "동일":
            continue
        nc, oc = p["new"], p["old"]
        if nc and ecs.load(new_doc).get(nc):
            before = ecs.load(data["old_doc"]).get(oc, {}).get("rows") if oc else None
            v = ecs.view_rows(new_doc, nc, before if before is not None else [])
            for r in v["rows"]:
                for f in ecs.FIELDS:
                    cell = r["fields"][f["key"]]
                    bc = cell.get("before") or {}
                    ch = (bc.get("code") or "", bc.get("name") or "") != (cell["code"], cell["name"])
                    ws3.append([nc, p["new_name"], r["label"], f["label"], bc.get("code"), bc.get("name"), cell["code"],
                                cell["name"], "변경" if ch else "", cell["status"],
                                ", ".join(str(e["page"]) for e in cell["evidence"] if e.get("page"))])
                    if ch:
                        for c in ws3[ws3.max_row]:
                            c.fill = fill
        else:
            for f in p["fields"]:
                ws3.append([nc or oc, p["new_name"] or p["old_name"], "(4단계 초안 없음)", f, "", "", "", "",
                            "수정 필요 추정", p["status"], ""])
    for wsx in (ws, ws2, ws3):
        for c in wsx[1]:
            c.font = Font(bold=True)
        for col in "ABCDEFGHIJK":
            wsx.column_dimensions[col].width = 22
    ws2.column_dimensions["H"].width = 70
    ws2.column_dimensions["I"].width = 70
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
