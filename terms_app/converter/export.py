"""엑셀(.xlsx)·JSON(zip) 내보내기. 선택한 담보만 내보낼 수 있습니다."""
from __future__ import annotations

import io
import json
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import pipeline

HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")
KIND_KO = {"article": "조", "para": "항", "item": "호", "sub": "목", "subsub": "세목", "box": "설명상자",
           "annex": "별표", "preface": "머리말", "part": "편", "chap": "관", "law": "법령"}


def _sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
    for r in rows:
        ws.append([("" if v is None else v) for v in r])
    for i, h in enumerate(header, 1):
        w = (widths or {}).get(h, 14)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


def load_all(doc_id):
    return {
        "cmap": pipeline.read_json(doc_id, "coverage_map.json"),
        "terms": pipeline.read_json(doc_id, "terms.json"),
        "tables": pipeline.read_json(doc_id, "tables.json"),
        "pages": pipeline.read_json(doc_id, "pages.json"),
        "report": pipeline.read_json(doc_id, "report.json"),
        "reviews": pipeline.read_json(doc_id, "reviews.json", {}) or {},
        "history": pipeline.read_json(doc_id, "history.json", []) or [],
        "approval": pipeline.read_json(doc_id, "approval.json", {}) or {},
        "meta": pipeline.read_json(doc_id, "meta.json", {}) or {},
    }


def excel(doc_id, cids=None):
    d = load_all(doc_id)
    covs = d["cmap"]["coverages"]
    sel = set(cids) if cids else {c["cid"] for c in covs}
    nodes = d["terms"]["nodes"]
    rep = d["report"]
    wb = Workbook()
    wb.remove(wb.active)

    def page_of(nid):
        return nodes[nid]["page_start"] if nid and nid in nodes else None

    rows = []
    for c in covs:
        if c["cid"] not in sel:
            continue
        k = c["key"]
        ann = "; ".join(f"{a['key']} {a['title']} (p.{a['page']})" for a in c["annexes"])
        rows.append([c["cid"], c["name"], c["start"], c["end"], c["type"],
                     nodes[k["payout"]]["key"] + " " + nodes[k["payout"]]["title"] if k["payout"] else "없음",
                     page_of(k["payout"]), ann, c["n_articles"], len(c["tables"]),
                     d["reviews"].get(c["cid"], {}).get("status", "미검토")])
    _sheet(wb, "담보별요약", ["담보ID", "담보명", "시작쪽", "끝쪽", "유형", "지급사유 조문", "지급사유 쪽", "별표(쪽)",
                         "조문 수", "표 수", "검토상태"], rows,
           {"담보명": 55, "지급사유 조문": 30, "별표(쪽)": 70})

    appr = d["approval"]
    info = [
        ["파일명", d["meta"].get("filename")],
        ["전체 쪽수", rep["page_count"]],
        ["목차 쪽", ", ".join(map(str, rep["toc"]["pages"]))],
        ["목차 항목 수", rep["toc_count"]],
        ["머리글로 찾은 담보 수", rep["header_coverages"]],
        ["차이", rep["count_diff"]],
        ["목차 쪽번호 - PDF 쪽번호 오프셋", rep["offset"]["offset"]],
        ["미배정 쪽 수", len(rep["unassigned_pages"])],
        ["목차·머리글 불일치 쪽 수", len(rep["mismatch_pages"])],
        ["지급사유 조문 없는 담보 수", len(rep["no_payout"])],
        ["조문 번호 이상 수", len(rep["numbering"])],
        ["참조 수(전체/미해결)", f"{rep['ref_total']} / {len(rep['unresolved_refs'])}"],
        ["표 수(전체/불완전)", f"{rep['table_total']} / {len(rep['table_incomplete'])}"],
        ["텍스트가 거의 없는 쪽 수", len(rep["low_text_pages"])],
        ["스캔형 쪽 수", rep["scanned_pages"]],
        ["변환 승인", "승인" if appr.get("approved") else "미승인"],
    ]
    _sheet(wb, "품질보고서", ["항목", "값"], info, {"항목": 34, "값": 40})

    _sheet(wb, "쪽배정", ["PDF쪽", "인쇄쪽", "머리글", "배정 담보", "배정 근거", "머리글 담보", "목차 범위 담보", "불일치"],
           [[p["p"], p["printed"], p["header"], p["cid"], p["source"], p["header_cid"], p["toc_cid"],
             "불일치" if p["mismatch"] else ""] for p in d["pages"]["pages"]], {"머리글": 50})
    _sheet(wb, "미배정쪽", ["PDF쪽", "첫 줄"], [[u["p"], u["first_line"]] for u in rep["unassigned_pages"]], {"첫 줄": 70})
    _sheet(wb, "목차머리글불일치", ["PDF쪽", "머리글", "머리글 담보", "목차 범위 담보", "최종 배정"],
           [[m["p"], m["header"], m["header_cid"], m["toc_cid"], m["assigned"]] for m in rep["mismatch_pages"]]
           + [[None, f"목차 시작쪽 {s['toc_start']} / 머리글 시작쪽 {s['header_start']}", s["cid"], s["name"], None]
              for s in rep["start_diff"]], {"머리글": 60})
    _sheet(wb, "지급사유없음", ["담보ID", "담보명", "유형"], [[x["cid"], x["name"], x["type"]] for x in rep["no_payout"]],
           {"담보명": 60})
    _sheet(wb, "조문번호이상", ["담보ID", "구분", "조문", "쪽", "내용"],
           [[x["cid"], x["kind"], x["key"], x["page"], x["detail"]] for x in rep["numbering"]]
           + [[x["cid"], "ID 중복", x["id"], x["page"], ""] for x in rep["duplicate_ids"]], {"내용": 60})
    _sheet(wb, "참조미해결", ["담보ID", "출처 ID", "쪽", "종류", "참조 표현", "사유", "문맥"],
           [[r["cid"], r["src"], r["src_page"], r["kind"], r["raw"], r["reason"], r["ctx"]]
            for r in rep["unresolved_refs"] if r["cid"] in sel], {"출처 ID": 50, "사유": 45, "문맥": 70})
    _sheet(wb, "표추출불완전", ["표ID", "담보ID", "쪽", "별표", "문제", "쪽 이미지"],
           [[t["id"], t["cid"], t["page"], t["annex"], "; ".join(t["problems"]), t["image"]]
            for t in rep["table_incomplete"]], {"문제": 40, "쪽 이미지": 30})
    _sheet(wb, "저텍스트쪽", ["PDF쪽", "글자 수", "담보ID"], [[x["p"], x["chars"], x["cid"]] for x in rep["low_text_pages"]])

    trows = []
    for c in covs:
        if c["cid"] not in sel:
            continue
        for nid in d["terms"]["orders"].get(c["cid"], []):
            n = nodes[nid]
            trows.append([nid, c["cid"], KIND_KO.get(n["kind"], n["kind"]), n["key"], n.get("title", ""),
                          n["page_start"], n["page_end"], n["text"]])
    ws = _sheet(wb, "조문", ["ID", "담보ID", "종류", "번호", "제목", "시작쪽", "끝쪽", "본문"], trows,
                {"ID": 60, "제목": 30, "본문": 90})
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=8):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet("표")
    ws.column_dimensions["A"].width = 14
    for t in d["tables"]["tables"]:
        if t["cid"] not in sel:
            continue
        ws.append([f"[{t['id']}]", f"담보 {t['cid']}", f"쪽 {t['page']}", f"소속 {t['owner'] or ''}",
                   f"제목줄 {t.get('caption', '')}", "완전" if t["complete"] else "표 추출 불완전: " + "; ".join(t["problems"])])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        for r in t["grid"]:
            ws.append([""] + [("〃" if v is None else v) for v in r])
        ws.append([])
    for col in range(2, 12):
        ws.column_dimensions[get_column_letter(col)].width = 30

    _sheet(wb, "수정이력", ["시각", "작업자", "구분", "내용"],
           [[h["at"], h["user"], h["what"], json.dumps(h["data"], ensure_ascii=False)] for h in d["history"]],
           {"내용": 100})
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def json_zip(doc_id, cids=None):
    d = load_all(doc_id)
    sel = set(cids) if cids else None
    cmap = dict(d["cmap"])
    terms = dict(d["terms"])
    tables = dict(d["tables"])
    pages = dict(d["pages"])
    if sel:
        cmap["coverages"] = [c for c in cmap["coverages"] if c["cid"] in sel]
        terms["nodes"] = {k: v for k, v in terms["nodes"].items() if v["cid"] in sel}
        terms["orders"] = {k: v for k, v in terms["orders"].items() if k in sel}
        terms["refs"] = [r for r in terms["refs"] if r["cid"] in sel]
        terms.pop("line_owner", None)
        tables["tables"] = [t for t in tables["tables"] if t["cid"] in sel]
        pages["pages"] = [p for p in pages["pages"] if p["cid"] in sel]
    for c in cmap["coverages"]:
        c["review"] = d["reviews"].get(c["cid"], {"status": "미검토"})
    bio = io.BytesIO()
    with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as z:
        for name, obj in (("coverage_map.json", cmap), ("terms.json", terms), ("tables.json", tables),
                          ("pages.json", pages), ("report.json", d["report"])):
            z.writestr(name, json.dumps(obj, ensure_ascii=False, indent=1))
    return bio.getvalue()
