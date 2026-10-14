"""내보내기: 담보 분류 엑셀(.xlsx)·JSON"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

COV_COLUMNS = [
    ("ID", "id", 8), ("대분류", "major", 8), ("중분류", "middle", 12), ("소분류(특약)", "minor", 34),
    ("담보명", "name", 34), ("자동 인식 이름", "auto_name", 30), ("급부 그룹", "group", 26),
    ("지급사유(지급기준표 원문)", "pay_reason", 60), ("지급금액(지급기준표 원문)", "pay_amount", 40),
    ("분류 근거", "reason_note", 50), ("구간", "section_title", 40), ("조문 범위", "range_text", 30),
    ("주요 조문", "key_text", 50), ("근거 위치", "source_text", 26), ("사용자 수정", "edited_text", 14),
]


def _row(it):
    r = dict(it)
    rng = it["range"]
    r["range_text"] = f"{rng['first']} ~ {rng['last']} ({rng['pages'][0]}~{rng['pages'][1]}쪽)"
    r["key_text"] = ", ".join(c.split("|", 1)[-1] for c in it["key_articles"])
    src = it["source"]
    r["source_text"] = (f"{src['appendix']} {src['table']} {src['row'] + 1}번째 급부 ({src['page']}쪽)"
                        if src["type"] == "지급기준표" else f"특약 전체 ({src['page']}쪽)")
    r["edited_text"] = " ".join(x for x, f in (("분류 이동", it["moved"]), ("이름 수정", it["renamed"])) if f)
    return r


def coverage_items(cov, selected_only=False):
    return [i for i in cov["items"] if i["selected"] or not selected_only]


def coverage_xlsx(cov, selected_only=False, title=""):
    wb = Workbook()
    ws = wb.active
    ws.title = "담보 분류"
    ws.append([c[0] for c in COV_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE6F0")
    for it in coverage_items(cov, selected_only):
        r = _row(it)
        ws.append([r.get(k, "") for _, k, _ in COV_COLUMNS])
    for i, (_, _, w) in enumerate(COV_COLUMNS, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    info = wb.create_sheet("정보")
    info.append(["약관", title])
    info.append(["범위", "선택한 담보만" if selected_only else "전체 담보"])
    info.append(["만든 시각", cov.get("built_from", {}).get("built_at", "")])
    info.append(["근거 승인 시각", cov.get("built_from", {}).get("approved_at", "")])
    info.append(["주의", "지급사유·지급금액은 지급기준표 원문을 옮긴 것이며 해석하지 않았습니다. 분류는 낱말 규칙에 따른 제안입니다."])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
